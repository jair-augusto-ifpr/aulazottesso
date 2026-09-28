"""Extração aprimorada de texto e blocos estruturados de arquivos anexados ao Material."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List

from docx import Document
from pypdf import PdfReader


def _decode_bytes(raw: bytes) -> str:
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _clean_text(text: str) -> str:
    """Limpa espaços extras sem remover números, datas, regras e pontuação essencial."""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


@dataclass
class ExtractedBlock:
    """Bloco estruturado extraído de um arquivo (com página ou seção real)."""

    text: str
    page_or_section: str = ""
    block_type: str = "text"  # 'text', 'table', 'heading'
    char_start: int = 0
    char_end: int = 0


@dataclass
class ExtractionResult:
    """Resultado detalhado da extração documental."""

    text: str
    blocks: List[ExtractedBlock] = field(default_factory=list)
    status: str = "pending"  # 'success', 'empty', 'scanned', 'unsupported', 'error'
    error_message: str = ""
    page_count: int = 0


def extract_document_blocks(uploaded_file) -> ExtractionResult:
    """Extrai blocos estruturados e texto de PDF/DOCX/TXT/MD preservando metadados de página/seção."""
    name = getattr(uploaded_file, "name", "") or ""
    suffix = Path(name).suffix.lower()

    try:
        uploaded_file.seek(0)
    except Exception:
        pass

    if suffix not in {".pdf", ".docx", ".txt", ".md"}:
        return ExtractionResult(
            text="",
            blocks=[],
            status="unsupported",
            error_message=f"Formato '{suffix}' não suportado. Formatos suportados: PDF, DOCX, TXT, MD.",
        )

    try:
        if suffix == ".pdf":
            reader = PdfReader(uploaded_file)
            page_count = len(reader.pages)
            blocks: list[ExtractedBlock] = []
            full_texts: list[str] = []
            char_cursor = 0

            for idx, page in enumerate(reader.pages, start=1):
                raw_page_text = page.extract_text() or ""
                cleaned_page = _clean_text(raw_page_text)
                if cleaned_page:
                    start = char_cursor
                    end = start + len(cleaned_page)
                    blocks.append(
                        ExtractedBlock(
                            text=cleaned_page,
                            page_or_section=f"Página {idx}",
                            block_type="text",
                            char_start=start,
                            char_end=end,
                        )
                    )
                    full_texts.append(cleaned_page)
                    char_cursor = end + 2  # para \n\n

            combined = "\n\n".join(full_texts).strip()

            # Detecção de possível documento escaneado (páginas existem, mas texto quase nulo)
            if page_count > 0 and len(combined) < 40:
                return ExtractionResult(
                    text=combined,
                    blocks=blocks,
                    status="scanned",
                    error_message=(
                        "O arquivo possui páginas, mas quase nenhum texto selecionável foi extraído. "
                        "Pode ser um documento escaneado. Insira o texto manualmente ou use OCR."
                    ),
                    page_count=page_count,
                )

            if not combined:
                return ExtractionResult(
                    text="",
                    blocks=[],
                    status="empty",
                    error_message="O arquivo PDF está vazio.",
                    page_count=page_count,
                )

            return ExtractionResult(
                text=combined,
                blocks=blocks,
                status="success",
                page_count=page_count,
            )

        elif suffix == ".docx":
            doc = Document(uploaded_file)
            blocks = []
            full_texts = []
            char_cursor = 0
            section_idx = 1
            table_idx = 1

            # Processar parágrafos
            current_para_group = []
            for p in doc.paragraphs:
                p_text = _clean_text(p.text)
                if not p_text:
                    continue
                current_para_group.append(p_text)
                # Agrupa a cada ~3 parágrafos ou cabeçalhos
                if len(current_para_group) >= 3 or p.style.name.startswith("Heading"):
                    group_text = "\n".join(current_para_group)
                    start = char_cursor
                    end = start + len(group_text)
                    blocks.append(
                        ExtractedBlock(
                            text=group_text,
                            page_or_section=f"Seção {section_idx}",
                            block_type="heading" if p.style.name.startswith("Heading") else "text",
                            char_start=start,
                            char_end=end,
                        )
                    )
                    full_texts.append(group_text)
                    char_cursor = end + 2
                    section_idx += 1
                    current_para_group = []

            if current_para_group:
                group_text = "\n".join(current_para_group)
                start = char_cursor
                end = start + len(group_text)
                blocks.append(
                    ExtractedBlock(
                        text=group_text,
                        page_or_section=f"Seção {section_idx}",
                        block_type="text",
                        char_start=start,
                        char_end=end,
                    )
                )
                full_texts.append(group_text)
                char_cursor = end + 2
                section_idx += 1

            # Processar tabelas preservando linhas e colunas
            for table in doc.tables:
                table_lines = []
                for row in table.rows:
                    cells = [_clean_text(c.text).replace("\n", " ") for c in row.cells]
                    # Desduplica células adjacentes em células mescladas
                    deduped_cells = []
                    for c in cells:
                        if not deduped_cells or c != deduped_cells[-1]:
                            deduped_cells.append(c)
                    if any(deduped_cells):
                        table_lines.append(" | ".join(deduped_cells))
                if table_lines:
                    tbl_text = "\n".join(table_lines)
                    start = char_cursor
                    end = start + len(tbl_text)
                    blocks.append(
                        ExtractedBlock(
                            text=tbl_text,
                            page_or_section=f"Tabela {table_idx}",
                            block_type="table",
                            char_start=start,
                            char_end=end,
                        )
                    )
                    full_texts.append(tbl_text)
                    char_cursor = end + 2
                    table_idx += 1

            combined = "\n\n".join(full_texts).strip()
            if not combined:
                return ExtractionResult(
                    text="",
                    blocks=[],
                    status="empty",
                    error_message="O arquivo DOCX está vazio ou sem texto legível.",
                )

            return ExtractionResult(
                text=combined,
                blocks=blocks,
                status="success",
            )

        elif suffix in {".txt", ".md"}:
            raw_content = _decode_bytes(uploaded_file.read())
            combined = _clean_text(raw_content)
            if not combined:
                return ExtractionResult(
                    text="",
                    blocks=[],
                    status="empty",
                    error_message="O arquivo de texto está vazio.",
                )

            # Divide em parágrafos preservando seções
            raw_sections = [s.strip() for s in re.split(r"\n\s*\n", combined) if s.strip()]
            blocks = []
            char_cursor = 0
            for idx, sec in enumerate(raw_sections, start=1):
                start = char_cursor
                end = start + len(sec)
                blocks.append(
                    ExtractedBlock(
                        text=sec,
                        page_or_section=f"Seção {idx}",
                        block_type="text",
                        char_start=start,
                        char_end=end,
                    )
                )
                char_cursor = end + 2

            return ExtractionResult(
                text=combined,
                blocks=blocks,
                status="success",
            )

    except Exception as exc:
        return ExtractionResult(
            text="",
            blocks=[],
            status="error",
            error_message=f"Erro durante a extração do arquivo: {exc}",
        )
    finally:
        try:
            uploaded_file.seek(0)
        except Exception:
            pass

    return ExtractionResult(text="", blocks=[], status="error", error_message="Falha desconhecida.")


def extract_text_from_upload(uploaded_file) -> str:
    """Extrai texto simples de PDF/DOCX/TXT/MD mantendo compatibilidade com callers existentes."""
    res = extract_document_blocks(uploaded_file)
    return res.text


def apply_material_text_extraction(material, *, prefer_file: bool = False) -> int:
    """Atualiza text_content e metadados de status a partir do arquivo.

    Retorna caracteres extraídos ou 0.
    """
    if not material.file:
        return 0

    extracted = (extract_text_from_upload(material.file) or "").strip()

    update_fields = ["updated_at"]
    if hasattr(material, "extraction_status"):
        if extracted:
            material.extraction_status = "success"
            material.extraction_error = ""
        else:
            material.extraction_status = "empty"
        update_fields.extend(["extraction_status", "extraction_error"])

    if not extracted:
        material.save(update_fields=update_fields)
        return 0

    current = (material.text_content or "").strip()
    should_replace = prefer_file or not current or len(extracted) > len(current)
    if not should_replace:
        material.save(update_fields=update_fields)
        return 0

    material.text_content = extracted
    update_fields.append("text_content")
    material.save(update_fields=update_fields)
    return len(extracted)
